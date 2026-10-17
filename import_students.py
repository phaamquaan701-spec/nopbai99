import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import re
import unicodedata
from pathlib import Path
from datetime import datetime, timedelta
import openpyxl

from app.config import DEFAULT_TIMEZONE, BASE_DIR
from app.database import get_db, init_db
from app.security import hash_password

def remove_vietnamese_accents(text: str) -> str:
    text = unicodedata.normalize("NFD", text)
    text = re.sub(r"[\u0300-\u036f]", "", text)
    text = text.replace("đ", "d").replace("Đ", "D")
    return text.lower()

def generate_username(full_name: str, stt: int) -> str:
    clean_name = remove_vietnamese_accents(full_name)
    parts = clean_name.split()
    if not parts:
        return f"hs{stt:02d}"
    last_name = parts[-1]  # Tên chính (ví dụ: an, bach, duong)
    # Lấy chữ cái đầu họ đệm (ví dụ: pq cho pham quoc)
    initials = "".join([p[0] for p in parts[:-1]])
    return f"{last_name}_{initials}{stt:02d}"

def import_students_from_excel(excel_path: str):
    init_db()
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # Ưu tiên sheet Dữ liệu gốc vì đầy đủ môn và khối thi
    sheet = wb["Dữ liệu gốc"] if "Dữ liệu gốc" in wb.sheetnames else wb.active
    
    # Headers
    headers = [str(c.value).strip() if c.value is not None else "" for c in sheet[1]]
    
    # Tìm vị trí các cột
    col_stt = headers.index("STT") if "STT" in headers else 0
    col_name = headers.index("Họ và tên") if "Họ và tên" in headers else 1
    col_khoi = headers.index("Khối thi") if "Khối thi" in headers else -1

    subject_cols = {}
    for idx, h in enumerate(headers):
        if h in ["Toán", "Lý", "Hóa", "Sinh", "Anh", "Văn", "Sử", "Tin"]:
            subject_cols[h] = idx

    students = []
    for row_idx in range(2, sheet.max_row + 1):
        cells = sheet[row_idx]
        vals = [str(c.value).strip() if c.value is not None else "" for c in cells]
        
        if len(vals) <= col_name:
            continue
        name = vals[col_name]
        if not name or name.lower() == "none" or name.startswith("Họ và tên"):
            continue

        stt_val = vals[col_stt] if col_stt < len(vals) and vals[col_stt].isdigit() else str(len(students) + 1)
        stt = int(stt_val)
        
        khoi = vals[col_khoi] if col_khoi != -1 and col_khoi < len(vals) else ""

        # Môn học đăng ký
        subs = []
        for sub_name, col_i in subject_cols.items():
            if col_i < len(vals) and vals[col_i].lower() in ["x", "✓", "1", "yes", "true"]:
                subs.append(sub_name)

        username = generate_username(name, stt)
        default_pwd = "123456"

        students.append({
            "stt": stt,
            "full_name": name,
            "username": username,
            "password": default_pwd,
            "khoi": khoi,
            "subjects": subs
        })

    print(f"Đã trích xuất thành công {len(students)} học sinh từ file Excel.")

    # Lưu vào Database
    now = datetime.now(DEFAULT_TIMEZONE)
    class_name = "Lớp 12 - Luyện Thi THPT"

    with get_db() as conn:
        cursor = conn.cursor()

        # 1. Tạo lớp học
        cursor.execute("SELECT id FROM classes WHERE name = ?", (class_name,))
        row = cursor.fetchone()
        if not row:
            cursor.execute("INSERT INTO classes (name, created_at) VALUES (?, ?)", (class_name, now.isoformat()))
            class_id = cursor.lastrowid
        else:
            class_id = row["id"]

        # 2. Lấy admin_id
        admin_row = conn.execute("SELECT id FROM users WHERE role = 'admin' LIMIT 1").fetchone()
        admin_id = admin_row["id"] if admin_row else 1

        # 3. Tạo tài khoản học sinh
        pwd_hash = hash_password("123456")
        added_count = 0
        updated_count = 0

        for st in students:
            cursor.execute("SELECT id FROM users WHERE username = ?", (st["username"],))
            existing = cursor.fetchone()
            if existing:
                cursor.execute(
                    "UPDATE users SET full_name = ?, class_id = ? WHERE id = ?",
                    (st["full_name"], class_id, existing["id"])
                )
                updated_count += 1
            else:
                cursor.execute(
                    """
                    INSERT INTO users (username, password_hash, full_name, role, class_id, status, created_at)
                    VALUES (?, ?, ?, 'student', ?, 'active', ?)
                    """,
                    (st["username"], pwd_hash, st["full_name"], class_id, now.isoformat())
                )
                added_count += 1

        # 4. Tạo các bài tập mẫu thực tế cho lớp này nếu chưa có
        cursor.execute("SELECT COUNT(*) as c FROM assignments WHERE class_id = ?", (class_id,))
        if cursor.fetchone()["c"] == 0:
            sample_assignments = [
                ("Bài tập Toán – Khảo sát hàm số & Vẽ đồ thị", "Toán học", 3),
                ("Bài tập Vật lý – Dao động & Sóng cơ", "Vật lý", 2),
                ("Bài tập Tiếng Anh – Chuyên đề Ngữ pháp & Từ vựng", "Tiếng Anh", 4),
                ("Bài tập Hóa học – Chuyên đề Este & Lipit", "Hóa học", 5)
            ]
            for title, sub, days in sample_assignments:
                s_at = (now - timedelta(days=1)).isoformat()
                d_at = (now + timedelta(days=days)).replace(hour=23, minute=59, second=0).isoformat()
                cursor.execute(
                    """
                    INSERT INTO assignments (
                        title, description, subject, class_id, created_by,
                        start_at, deadline, allow_resubmit, allow_late_submission,
                        resubmit_after_deadline, submission_grace_period, max_images, max_file_size,
                        status, created_at, updated_at
                    )
                    VALUES (?, 'Học sinh làm bài ra giấy, chụp ảnh rõ nét và nộp trước deadline.', ?, ?, ?, ?, ?, 1, 1, 1, 300, 10, 10485760, 'open', ?, ?)
                    """,
                    (title, sub, class_id, admin_id, s_at, d_at, now.isoformat(), now.isoformat())
                )

    print(f"Cập nhật Database: Thêm mới {added_count} học sinh, cập nhật {updated_count} học sinh vào '{class_name}'.")

    # 5. Xuất file Excel bàn giao tài khoản cho giáo viên
    out_wb = openpyxl.Workbook()
    out_sheet = out_wb.active
    out_sheet.title = "Danh sách tài khoản"

    out_sheet.append(["STT", "Họ và tên", "Tên đăng nhập (Username)", "Mật khẩu mặc định", "Lớp học", "Khối thi", "Các môn đăng ký"])
    for st in students:
        out_sheet.append([
            st["stt"],
            st["full_name"],
            st["username"],
            st["password"],
            class_name,
            st["khoi"],
            ", ".join(st["subjects"])
        ])

    out_path = Path(r"c:\Users\quanp\Documents\web\Danh_sach_tai_khoan_hoc_sinh.xlsx")
    out_wb.save(str(out_path))
    print(f"Đã xuất file tài khoản bàn giao tại: {out_path}")

if __name__ == "__main__":
    excel_file = r"c:\Users\quanp\Documents\web\Danh_sach_hoc_sinh_dang_ky_mon.xlsx"
    import_students_from_excel(excel_file)
